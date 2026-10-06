"""Export primer QC results to CSV / XLSX / HTML."""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import List, Dict, Any

from .primer_analyzer import PrimerResult, PrimerQCStatus, QCConfig

logger = logging.getLogger(__name__)


def export_primer_results(
    results: List[PrimerResult],
    qc_statuses: List[PrimerQCStatus],
    config: QCConfig,
    errors: List[Dict[str, str]],
    output_path: Path,
) -> None:
    """Export all results based on output file extension."""
    suffix = output_path.suffix.lower()

    if suffix == ".csv":
        _export_csv(results, qc_statuses, output_path)
    elif suffix == ".html":
        _export_html(results, qc_statuses, output_path)
    else:
        # Default: XLSX
        _export_xlsx(results, qc_statuses, config, errors, output_path)
        # Also export companion CSV
        csv_path = output_path.with_suffix(".csv")
        _export_csv(results, qc_statuses, csv_path)

    logger.info("Exported %d primer results to %s", len(results), output_path)


def _export_csv(
    results: List[PrimerResult],
    qc_statuses: List[PrimerQCStatus],
    path: Path,
) -> None:
    """Export primary results to CSV."""
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([
            "Primer_ID", "Length", "GC_Percent", "Tm_C", "Tm_Method",
            "A_Count", "C_Count", "G_Count", "T_Count",
            "Ambiguous_Count", "Homopolymer", "GC_Clamp",
            "Self_Complementarity", "Hairpin_Risk",
            "QC_Status", "Warnings",
        ])
        for res, qc in zip(results, qc_statuses):
            writer.writerow([
                res.primer_id, res.length, res.gc_percent, res.tm, res.tm_method,
                res.base_counts.get("A", 0), res.base_counts.get("C", 0),
                res.base_counts.get("G", 0), res.base_counts.get("T", 0),
                res.ambiguous_count, res.homopolymer,
                res.gc_clamp, res.self_complementarity, res.hairpin_risk,
                qc.status, "; ".join(qc.warnings),
            ])


def _export_xlsx(
    results: List[PrimerResult],
    qc_statuses: List[PrimerQCStatus],
    config: QCConfig,
    errors: List[Dict[str, str]],
    path: Path,
) -> None:
    """Export to multi-sheet XLSX."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    hdr_font = Font(bold=True)
    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    # Summary
    ws = wb.active
    ws.title = "Summary"
    summary = [
        ("Metric", "Value"),
        ("Total primers", len(results)),
        ("PASS", sum(1 for q in qc_statuses if q.status == "PASS")),
        ("REVIEW", sum(1 for q in qc_statuses if q.status == "REVIEW")),
        ("WARNING", sum(1 for q in qc_statuses if q.status == "WARNING")),
        ("INVALID", sum(1 for q in qc_statuses if q.status == "INVALID")),
        ("Errors", len(errors)),
    ]
    for row in summary:
        ws.append(row)
    for cell in ws[1]:
        cell.font = hdr_font

    # Primer Results
    ws2 = wb.create_sheet("Primer Results")
    headers = ["Primer_ID", "Sequence", "Length", "GC%", "Tm(°C)", "Tm_Method",
               "Homopolymer", "Self-Comp", "Hairpin", "QC_Status", "Warnings"]
    ws2.append(headers)
    for c in ws2[1]:
        c.font = hdr_font
    for res, qc in zip(results, qc_statuses):
        ws2.append([
            res.primer_id, res.sequence, res.length, res.gc_percent,
            res.tm, res.tm_method, res.homopolymer,
            res.self_complementarity, res.hairpin_risk,
            qc.status, "; ".join(qc.warnings),
        ])

    # QC Rules
    ws3 = wb.create_sheet("QC Rules")
    rules = [
        ("Rule", "Value"),
        ("Length preferred", f"{config.min_length}-{config.max_length}"),
        ("GC% preferred", f"{config.gc_min}-{config.gc_max}%"),
        ("Tm preferred", f"{config.tm_min}-{config.tm_max}°C"),
        ("Homopolymer threshold", str(config.homopolymer_max)),
        ("Salt concentration", f"{config.salt_conc} mM"),
        ("Mg2+ concentration", f"{config.mg_conc} mM"),
        ("Primer concentration", f"{config.primer_conc} µM"),
        ("Tm method", config.tm_method),
    ]
    for row in rules:
        ws3.append(row)

    # Methods
    ws4 = wb.create_sheet("Methods")
    methods = [
        ("Method", "Description"),
        ("Tm (basic)", "4*(G+C) + 2*(A+T), salt-corrected"),
        ("Tm (nearest_neighbor)", "SantaLucia 1998 unified NN parameters"),
        ("GC%", "(G+C)/(A+C+G+T) × 100"),
        ("Hairpin", "Computational self-complementarity heuristic"),
        ("Self-complementarity", "Overlapping reverse-complement stretch detection"),
        ("", ""),
        ("Disclaimer", "Computational QC estimates — does not guarantee PCR success, "
                       "no primer-dimer, or no hairpin formation."),
        ("Source", "BioSyn Primer Batch QC v1.0.0 — https://biosyn-inc.com/"),
    ]
    for row in methods:
        ws4.append(row)

    wb.save(path)


def _export_html(
    results: List[PrimerResult],
    qc_statuses: List[PrimerQCStatus],
    path: Path,
) -> None:
    """Export standalone HTML report."""
    rows = ""
    for res, qc in zip(results, qc_statuses):
        color = {"PASS": "#28a745", "REVIEW": "#ffc107",
                 "WARNING": "#fd7e14", "INVALID": "#dc3545"}.get(qc.status, "#6c757d")
        rows += f"""<tr><td>{res.primer_id}</td><td><code>{res.sequence}</code></td>
        <td>{res.length}</td><td>{res.gc_percent:.1f}%</td><td>{res.tm:.1f}°C</td>
        <td style="color:{color};font-weight:bold">{qc.status}</td>
        <td>{'; '.join(qc.warnings)}</td></tr>"""

    html = f"""<!DOCTYPE html><html lang="en"><head><meta charset="UTF-8">
<title>BioSyn Primer Batch QC Report</title>
<style>
body{{font-family:system-ui;margin:2rem;background:#f8f9fa}}
table{{border-collapse:collapse;width:100%;background:#fff}}
th,td{{padding:.5rem 1rem;border:1px solid #dee2e6;text-align:left}}
th{{background:#e9ecef}}
</style></head><body>
<h1>BioSyn Primer Batch QC Report</h1>
<p>By <a href="https://biosyn-inc.com/">BioSyn Inc.</a> — computational primer QC &amp; decision support.</p>
<table><tr><th>ID</th><th>Sequence</th><th>Len</th><th>GC%</th><th>Tm</th>
<th>QC</th><th>Warnings</th></tr>{rows}</table>
<p style="margin-top:2rem;font-size:.85rem;color:#6c757d">
Computational QC estimates. Not a guarantee of PCR success.
See <a href="https://biosyn-inc.com/tools/primer-design-qc">BioSyn Primer Design & QC Tools</a> for more tools.</p>
</body></html>"""
    path.write_text(html, encoding="utf-8")
